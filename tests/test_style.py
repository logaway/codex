import json
import os

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw
from pptx import Presentation

from deckforge import get_deck_tools, render_deck
from deckforge.style import contrast, extract_style, parse_color
from test_deckforge import DECK


def slide_png(path, bg, ink, accent, left=96):
    """画一张 1280×720 的「参考幻灯片」：标题条、强调色条、几行正文。"""
    im = Image.new("RGB", (1280, 720), bg)
    d = ImageDraw.Draw(im)
    d.rectangle([left, 80, left + 40, 88], fill=accent)
    d.rectangle([left, 120, left + 700, 170], fill=ink)
    for i in range(5):
        d.rectangle([left, 260 + i * 50, left + 900, 272 + i * 50], fill=ink)
    d.rectangle([left, 560, left + 300, 620], fill=accent)
    im.save(path)
    return str(path)


def test_image_light_reference(tmp_path):
    png = slide_png(tmp_path / "ref.png", "#FAF7F2", "#1A1A1A", "#0B6E4F", left=128)
    rep = extract_style(png, str(tmp_path / "styles"))
    c = rep["colors"]
    assert c["bg"] == "#FAF7F2"
    assert contrast(c["accent"], "#0B6E4F") < 1.3          # 取到的是那个绿色
    assert rep["based_on"] == "editorial-report"
    assert rep["margin_x"] == 128
    assert {"bg", "ink", "accent", "margin.x"} <= set(rep["from_file"])


def test_small_accent_text_is_found(tmp_path):
    """真实幻灯片里强调色常常只是一行小字眉标，只占千分之一的像素。"""
    im = Image.new("RGB", (1280, 720), "#FFFFFF")
    d = ImageDraw.Draw(im)
    d.text((96, 60), "Q3 REVIEW  OVERVIEW", fill="#B4351E")
    d.rectangle([96, 100, 800, 150], fill="#0A0A0A")
    im.save(tmp_path / "small.png")
    rep = extract_style(str(tmp_path / "small.png"), str(tmp_path / "styles"))
    assert "accent" in rep["from_file"]
    assert contrast(rep["colors"]["accent"], "#B4351E") < 1.6


def test_image_dark_reference_picks_dark_base(tmp_path):
    png = slide_png(tmp_path / "dark.png", "#0E1116", "#EDEDED", "#FF7A00")
    rep = extract_style(png, str(tmp_path / "styles"))
    assert rep["based_on"] == "technical-deck"
    for role, need in (("ink", 7), ("body", 4.5), ("muted", 3), ("accent", 3)):
        assert contrast(rep["colors"][role], rep["colors"]["bg"]) >= need, role


def test_low_contrast_brand_color_is_fixed_and_reported(tmp_path):
    png = slide_png(tmp_path / "yellow.png", "#FFFFFF", "#111111", "#FFE14D")
    rep = extract_style(png, str(tmp_path / "styles"))
    assert contrast(rep["colors"]["accent"], "#FFFFFF") >= 3
    assert any("accent" in a for a in rep["adjustments"])


def test_non_slide_image_keeps_base_margins(tmp_path):
    im = Image.new("RGB", (400, 400), "#FFFFFF")
    ImageDraw.Draw(im).ellipse([100, 100, 300, 300], fill="#C8102E")  # 一个 logo
    im.save(tmp_path / "logo.png")
    rep = extract_style(str(tmp_path / "logo.png"), str(tmp_path / "styles"))
    assert "margin.x" not in rep["from_file"]
    assert contrast(rep["colors"]["accent"], "#C8102E") < 1.3


def test_excel_spec_sheet(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["名称", "色值", "说明"])
    ws.append(["品牌主色", "#C8102E", "用于强调"])
    ws.append(["背景色", "RGB(255, 255, 255)"])
    ws.append(["正文颜色", "333333"])
    ws.append(["标题字体", "思源黑体 Heavy"])
    ws.append(["正文字体", "思源黑体"])
    ws.append(["辅助色"])
    ws["B7"].fill = PatternFill("solid", fgColor="FF8C8C8C")    # 色块单元格
    ws.append(["销量", 123456])                                # 数字不能当颜色
    wb.save(tmp_path / "brand.xlsx")

    rep = extract_style(str(tmp_path / "brand.xlsx"), str(tmp_path / "styles"), name="Acme Brand")
    assert rep["name"] == "acme-brand"
    c = rep["colors"]
    assert c["accent"] == "#C8102E" and c["bg"] == "#FFFFFF"
    assert c["body"] == "#333333" and c["muted"] == "#8C8C8C"
    assert rep["fonts"]["heading"] == "思源黑体 Heavy"
    assert rep["fonts"]["sans"] == "思源黑体"
    assert any("思源黑体" in n for n in rep["notes"])
    saved = json.load(open(tmp_path / "styles" / "acme-brand.json", encoding="utf-8"))
    assert saved["type"]["h2"]["font"] == "heading"


def test_excel_falls_back_to_cell_styles(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["区域", "营收", "增速"])
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="FF00539B")
        c.font = Font(color="FFFFFFFF")
    for r in (["华东", 120, "3%"], ["华南", 98, "5%"]):
        ws.append(r)
    wb.save(tmp_path / "report.xlsx")
    rep = extract_style(str(tmp_path / "report.xlsx"), str(tmp_path / "styles"))
    assert rep["colors"]["accent"] == "#00539B"
    assert "单元格" in rep["from_file"]["accent"]


def test_default_office_theme_is_ignored(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["a", 1])
    wb.save(tmp_path / "plain.xlsx")
    rep = extract_style(str(tmp_path / "plain.xlsx"), str(tmp_path / "styles"))
    assert rep["from_file"] == {}
    assert rep["notes"]


def test_parse_color():
    assert parse_color("#b4351e") == "#B4351E"
    assert parse_color("色值 #B4351E") == "#B4351E"
    assert parse_color("180, 53, 30") == "#B4351E"
    assert parse_color("思源黑体") is None


def test_unsupported_and_builtin_names(tmp_path):
    (tmp_path / "x.xls").write_bytes(b"")
    with pytest.raises(ValueError, match="xlsx"):
        extract_style(str(tmp_path / "x.xls"), str(tmp_path / "s"))
    png = slide_png(tmp_path / "r.png", "#FFFFFF", "#111111", "#C8102E")
    assert extract_style(png, str(tmp_path / "s"), name="exec-brief")["name"] != "exec-brief"


def test_render_with_extracted_style(tmp_path):
    png = slide_png(tmp_path / "ref.png", "#0E1116", "#EDEDED", "#FF7A00")
    styles = str(tmp_path / "styles")
    name = extract_style(png, styles)["name"]
    deck = dict(DECK, meta=dict(DECK["meta"], contract=name))
    res = render_deck(deck, out_dir=str(tmp_path / "out"), contracts_dir=styles)
    assert res["ok"], res
    prs = Presentation(res["pptx"])
    bg = prs.slides[0].background.fill.fore_color.rgb
    assert str(bg) == "0E1116"
    from pptx.oxml.ns import qn
    run = next(sh for sh in prs.slides[0].shapes if sh.has_text_frame).text_frame.paragraphs[0].runs[0]
    assert run.font._rPr.find(qn("a:ea")).get("typeface") == run.font.name


def test_style_tool_flow(tmp_path):
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    png = slide_png(uploads / "ref.png", "#FFFFFF", "#111111", "#C8102E")
    tools = {t.name: t for t in get_deck_tools(out_dir=str(tmp_path / "out"),
                                               uploads_dir=str(uploads))}
    rep = json.loads(tools["extract_style"].invoke({"file_path": png}))
    assert rep["ok"] and rep["name"] in rep["next"]
    assert rep["name"] in tools["list_deck_contracts"].invoke({})

    deck = dict(DECK, meta={"contract": rep["name"]})
    out = json.loads(tools["render_deck"].invoke({"deck": deck}))
    assert out["ok"] and out["contract"] == rep["name"]

    # 上传目录以外的文件一律拒绝
    secret = tmp_path / "secret.png"
    slide_png(secret, "#FFFFFF", "#111111", "#C8102E")
    rep = json.loads(tools["extract_style"].invoke({"file_path": str(secret)}))
    assert not rep["ok"]
    # 损坏的文件返回原因而不是抛异常
    (uploads / "broken.png").write_bytes(b"not an image")
    rep = json.loads(tools["extract_style"].invoke({"file_path": str(uploads / "broken.png")}))
    assert not rep["ok"] and "error" in rep
