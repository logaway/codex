import json
import os

import pytest
from pptx import Presentation

from deckforge import check_deck, create_deck_author, get_deck_tools, render_deck
from deckforge.chain import FEWSHOT_OUT
from deckforge.schema import Deck

DECK = {
    "meta": {"contract": "exec-brief", "org": "NORTHWIND", "date": "Q3 2026"},
    "pages": FEWSHOT_OUT["pages"] + [
        {"shape": "timeline", "title": "分三步收紧折扣，每步都可回滚",
         "steps": [{"when": "10月", "heading": "收回审批权", "body": "超过 15% 的折扣需总监审批。"},
                   {"when": "11月", "heading": "试点华东", "body": "观察签约周期变化。"},
                   {"when": "12月", "heading": "全面推行", "body": "按试点数据调整阈值。"}]},
        {"shape": "table", "title": "华东在三项指标上落后于其他大区",
         "columns": ["指标", "华东", "其他", "差距"],
         "rows": [["获客成本", "¥3.2k", "¥2.1k", "+52%"], ["续约率", "81%", "90%", "-9pt"],
                  ["签约周期", "46 天", "31 天", "+15 天"]]},
        {"shape": "pipeline", "title": "折扣治理处于方案设计阶段",
         "stages": [{"label": "诊断", "state": "done", "note": "已定位折扣失控的三个来源。"},
                    {"label": "设计", "state": "current", "note": "审批阈值与例外流程待定。"},
                    {"label": "试点", "note": "11 月在华东启动。"}]},
    ],
}


def test_render_produces_editable_pptx(tmp_path):
    res = render_deck(DECK, out_dir=str(tmp_path))
    assert res["ok"], res
    assert res["tally"]["P0"] == 0
    prs = Presentation(res["pptx"])
    assert len(prs.slides) == len(DECK["pages"])
    texts = [s.text_frame.text for sl in prs.slides for s in sl.shapes if s.has_text_frame]
    assert "104%" in texts            # 原生文本框，不是图片
    assert res["avoid"]


def test_invalid_deck_returns_specific_errors(tmp_path):
    _, errors = check_deck({"pages": [{"shape": "metrics", "title": "x", "items": []}]})
    assert any("第 1 頁必須是 title" in e for e in errors)
    res = render_deck({"pages": [{"shape": "nope"}]}, out_dir=str(tmp_path))
    assert not res["ok"] and "shape='nope'" in res["errors"][0]
    res = render_deck(DECK, out_dir=str(tmp_path), contract="missing")
    assert not res["ok"] and "missing" in res["errors"][0]


def test_label_titles_are_flagged(tmp_path):
    deck = json.loads(json.dumps(DECK))
    deck["pages"][1]["title"] = "市场分析"
    res = render_deck(deck, out_dir=str(tmp_path))
    assert any("像標籤" in w for w in res["warnings"])


def test_langchain_tools(tmp_path):
    tools = {t.name: t for t in get_deck_tools(out_dir=str(tmp_path))}
    assert set(tools) == {"list_deck_contracts", "extract_style",
                          "validate_deck", "render_deck"}
    names = [c["name"] for c in json.loads(tools["list_deck_contracts"].invoke({}))]
    assert "exec-brief" in names

    # 模拟 agent 发起的 tool call
    msg = tools["render_deck"].invoke(
        {"type": "tool_call", "id": "c1", "name": "render_deck", "args": {"deck": DECK}})
    out = json.loads(msg.content)
    assert out["ok"] and os.path.exists(out["pptx"])

    # 结构不合法时把错误返回给 agent，而不是抛异常
    msg = tools["render_deck"].invoke(
        {"type": "tool_call", "id": "c2", "name": "render_deck",
         "args": {"deck": {"pages": [{"shape": "bogus"}]}}})
    assert "deck.pages.0" in msg.content and "bogus" in msg.content

    # 给模型的 JSON schema 能正常生成，且不含任何几何字段
    schema = json.dumps(tools["render_deck"].tool_call_schema.model_json_schema())
    for geom in ('"x"', '"width"', '"font_size"'):
        assert geom not in schema


class FakeLLM:
    """按顺序返回预设输出的结构化模型替身，用来验证修复循环。"""

    def __init__(self, outputs):
        self.outputs, self.calls = list(outputs), []

    def with_structured_output(self, schema):
        llm = self

        class Runner:
            def invoke(self, msgs):
                llm.calls.append(msgs)
                return schema.model_validate(llm.outputs.pop(0))
        return Runner()


def test_author_repairs_with_specific_feedback():
    too_long = json.loads(json.dumps(DECK))
    too_long["pages"][1]["items"] = too_long["pages"][1]["items"] * 3   # 6 项，coerce 会截到 4
    too_long["pages"][2]["sides"][0]["rows"] = ["一"]                   # 只有 1 条，必须修
    llm = FakeLLM([too_long, DECK])
    deck = create_deck_author(llm, contract="exec-brief").invoke("Q3 复盘")
    assert deck["meta"]["contract"] == "exec-brief"
    assert len(llm.calls) == 2
    assert "rows 需 2–5 條" in llm.calls[1][-1].content


def test_author_gives_up_after_rounds():
    bad = {"pages": [{"shape": "points", "title": "只有一页且不是封面的要点页",
                      "items": [{"heading": "a"}, {"heading": "b"}]}]}
    with pytest.raises(ValueError):
        create_deck_author(FakeLLM([bad] * 2), rounds=2).invoke("x")


def test_wrapped_paragraph_lines_stay_together(tmp_path):
    """stretch 拉开层间距时，同一段折出的行必须保持正常行距。"""
    deck = {"meta": {"contract": "technical-deck"}, "pages": [
        {"shape": "title", "title": "封面标题写结论"},
        {"shape": "points", "title": "三项能力覆盖报销从录入到入账的全流程",
         "items": [{"heading": "拍照识别发票", "body": "员工拍照即可自动识别发票信息，免去手工录入，" * 2},
                   {"heading": "审批流可配置", "body": "按规则配置审批链。"},
                   {"heading": "对接主流 ERP", "body": "单据同步入账。"}]}]}
    res = render_deck(deck, out_dir=str(tmp_path))
    assert res["ok"], res
    dl = json.load(open(tmp_path / "pages" / "display-list.json", encoding="utf-8"))
    body = [o for o in dl["pages"][1]["ops"] if o["op"] == "text" and o["size"] == 17
            and o["x"] == min(p["x"] for p in dl["pages"][1]["ops"] if p["op"] == "text")]
    ys = sorted(o["y"] for o in body)
    gaps = {b - a for a, b in zip(ys, ys[1:])}
    assert len(ys) >= 2 and len(gaps) == 1 and gaps.pop() <= 17 * 1.8
