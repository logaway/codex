#!/usr/bin/env python3
"""
topptx.py — 把 display list 發射成原生可編輯的 PPTX。

不是把 SVG 拍平成圖片。每個文字都是真的 text frame、每個矩形都是真的
autoshape，打開 PowerPoint 可以直接改字、換色、調位置。

之所以能做到，是因為佈局階段已經把幾何算成了確定數值 —— 走提示詞路線的
方案拿不到結構化坐標，只能光柵化。

用法：
    python3 build.py deck.json --out ./pages
    python3 topptx.py ./pages/display-list.json --out deck.pptx
"""

import argparse
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# SVG 用 px（96dpi），PPTX 用 EMU。1px = 9525 EMU
PX = 9525


def emu(px):
    return Emu(int(round(px * PX)))


def rgb(hex_str):
    h = hex_str.lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_rect(slide, op):
    from pptx.enum.shapes import MSO_SHAPE
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if op.get("rx") else MSO_SHAPE.RECTANGLE
    # 極細的線（1px 分隔線）用矩形表達，圓角會失真，強制直角
    if op["h"] <= 2 or op["w"] <= 2:
        shape = MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, emu(op["x"]), emu(op["y"]),
                               emu(max(op["w"], 1)), emu(max(op["h"], 1)))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(op["fill"])
    s.line.fill.background()
    s.line.width = 0
    s.shadow.inherit = False
    # python-pptx 建立形狀時會附一個 <p:style>，引用佈景主題的線條與效果，
    # 會蓋過上面的設定 —— 髮絲線因此被渲染成帶漸層的粗灰條。整個移除。
    el = s._element
    style = el.find(qn("p:style"))
    if style is not None:
        el.remove(style)
    return s


def add_text(slide, op, canvas_w):
    """SVG 的 y 是基線，PPTX 的 y 是文字框頂端 —— 這個換算做錯會整體偏移。

    寬度不能用估算值。我們的 measure() 是給 CJK 字體校準的，PowerPoint 換用
    Inter/Arial 後實際更寬，框就撐破、在錯誤的地方換行（"104%" 斷成兩行）。
    改成給到畫布邊緣：文字框寬一點沒有任何視覺代價，換行卻是致命的。"""
    size = op["size"]
    box_h = size * 1.7
    top = op["y"] - size * 0.9
    left = op["x"]
    align = PP_ALIGN.LEFT
    if op["anchor"] == "end":
        left = 0
        width = op["x"]
        align = PP_ALIGN.RIGHT
    elif op["anchor"] == "middle":
        half = min(op["x"], canvas_w - op["x"])
        left = op["x"] - half
        width = half * 2
        align = PP_ALIGN.CENTER
    else:
        width = canvas_w - left      # 一路給到右邊緣，保證不換行

    tb = slide.shapes.add_textbox(emu(left), emu(top), emu(width), emu(box_h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP

    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = op["s"]
    f = run.font
    f.name = op["font"]
    f.size = Pt(size * 0.75)               # px → pt
    f.bold = int(op.get("weight", 400)) >= 600
    f.color.rgb = rgb(op["fill"])
    if op.get("tracking"):
        # python-pptx 沒有字距 API，直接寫 OOXML 屬性（spc 單位為 1/100 pt）
        run.font._rPr.set("spc", str(int(op["tracking"] * 0.75 * 100)))
    return tb


def convert(dl_path, out_path):
    dl = json.load(open(dl_path, encoding="utf-8"))
    W, H = dl["canvas"]["w"], dl["canvas"]["h"]

    prs = Presentation()
    prs.slide_width = emu(W)
    prs.slide_height = emu(H)
    blank = prs.slide_layouts[6]

    counts = {"text": 0, "rect": 0}
    for page in dl["pages"]:
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(dl["bg"])
        for op in page["ops"]:
            if op["op"] == "rect":
                add_rect(slide, op)
                counts["rect"] += 1
            elif op["op"] == "text":
                add_text(slide, op, W)
                counts["text"] += 1

    prs.save(out_path)
    print(f"{out_path} · {len(dl['pages'])} 頁 · "
          f"{counts['text']} 個文字框 + {counts['rect']} 個形狀，全部原生可編輯")
    return out_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("display_list")
    ap.add_argument("--out", default="deck.pptx")
    a = ap.parse_args()
    convert(a.display_list, a.out)
